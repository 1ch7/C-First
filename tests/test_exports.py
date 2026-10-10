import copy
import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from PIL import Image, ImageChops
import openpyxl
from pypdf import PdfReader
import server
from src.exports import budget_xlsx_bytes, proposal_pdf_bytes


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pack = json.loads((Path(__file__).resolve().parents[1] / 'data' / 'recording.json').read_text(encoding='utf-8'))['pack']

    def test_canonical_workbook_figures_chart_and_full_rows(self):
        original = copy.deepcopy(self.pack)
        raw = budget_xlsx_bytes(self.pack)
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        self.assertEqual(wb.sheetnames, ['Budget summary', 'Full breakdown', 'Category totals', 'Notes and source', 'Budget charts'])
        ws = wb['Budget summary']
        self.assertEqual([ws.cell(i, 2).value for i in range(2, 7)], [2000, 2000, 240, 600, 1400])
        details = wb['Full breakdown']
        self.assertEqual(details.max_row, len(self.pack['budget']['items']) + 1)
        self.assertEqual(sum(details.cell(i, 1).value for i in range(2, details.max_row + 1)), 2000)
        categories = wb['Category totals']
        self.assertEqual(sum(categories.cell(i, 2).value for i in range(2, categories.max_row + 1)), 2000)
        self.assertEqual(len(wb['Budget charts']._charts), 2)
        self.assertEqual(len(ws._images), 1)
        self.assertIn('Budget charts', ws['B9'].hyperlink.target)
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        with ZipFile(io.BytesIO(raw)) as archive:
            chart = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            cached_labels = [node.text for node in chart.findall('.//c:cat/c:strRef/c:strCache/c:pt/c:v', ns)]
            cached_amounts = [float(node.text) for node in chart.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v', ns)]
            self.assertEqual(cached_labels, [categories.cell(i, 1).value for i in range(2, categories.max_row + 1)])
            self.assertEqual(sum(cached_amounts), 2000)
            self.assertEqual(chart.find('.//c:cat/c:strRef/c:f', ns).text, "'Category totals'!$A$2:$A$9")
            self.assertEqual(chart.find('.//c:val/c:numRef/c:f', ns).text, "'Category totals'!$B$2:$B$9")
            # Category scale must not be assigned the monetary budget ceiling.
            self.assertIsNone(chart.find('.//c:catAx/c:scaling/c:max', ns))
            self.assertAlmostEqual(float(chart.find('.//c:valAx/c:scaling/c:max', ns).get('val')), 440 * 1.18)
            funding_chart = ET.fromstring(archive.read('xl/charts/chart2.xml'))
            funding_amounts = [float(n.text) for n in funding_chart.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v', ns)]
            self.assertEqual(funding_amounts, [2000, 2000, 240, 600, 1400])
            png = archive.read('xl/media/image1.png')
            image = Image.open(io.BytesIO(png))
            self.assertEqual(image.width, 1280)
            self.assertGreater(image.height, 800)
            self.assertIsNotNone(ImageChops.difference(image, Image.new('RGB', image.size, 'white')).getbbox())
        self.assertEqual(self.pack, original)

    def test_exports_use_custom_pack_not_hardcoded_demo(self):
        pack = copy.deepcopy(self.pack)
        for key in ['total', 'limit', 'contingency', 'gap', 'organiserFunds']:
            pack['budget'][key] *= 2
            pack['budget'][key + 'Cents'] *= 2
        for item in pack['budget']['items']:
            item['amountRM'] *= 2
            item['amountCents'] *= 2
        wb = openpyxl.load_workbook(io.BytesIO(budget_xlsx_bytes(pack)), data_only=True)
        self.assertEqual([wb['Budget summary'].cell(i, 2).value for i in range(2, 7)], [4000, 4000, 480, 1200, 2800])
        pdf = PdfReader(io.BytesIO(proposal_pdf_bytes(pack)))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertIn('RM4,000.00', text)
        self.assertIn('RM1,200.00', text)

    def test_missing_values_and_formula_injection(self):
        pack = {'budget': {'items': [{'category': '=HYPERLINK("https://example.com")', 'status': '+command', '=1+1': 'untrusted header'}]}}
        wb = openpyxl.load_workbook(io.BytesIO(budget_xlsx_bytes(pack)), data_only=False)
        self.assertIsNone(wb['Budget summary']['B2'].value)
        for ws in wb:
            for row in ws:
                self.assertTrue(all(cell.data_type != 'f' for cell in row))

    def test_chart_cache_preserves_missing_values_and_indices(self):
        pack = {'budget': {'limitCents': 100000, 'items': [
            {'category': 'Known', 'amountCents': 12550},
            {'category': 'Unknown'},
            {'category': 'Recorded zero', 'amountCents': 0},
        ]}}
        raw = budget_xlsx_bytes(pack)
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        with ZipFile(io.BytesIO(raw)) as archive:
            chart = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            cache = chart.find('.//c:val/c:numRef/c:numCache', ns)
            self.assertEqual(cache.find('c:ptCount', ns).get('val'), '3')
            points = {p.get('idx'): float(p.find('c:v', ns).text) for p in cache.findall('c:pt', ns)}
            self.assertEqual(points, {'0': 125.5, '2': 0.0})
            self.assertEqual(len(chart.findall('.//c:cat/c:strRef/c:strCache/c:pt', ns)), 3)
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        self.assertIsNone(wb['Category totals']['B3'].value)

    def test_empty_budget_has_visible_preview_and_no_fake_category(self):
        raw = budget_xlsx_bytes({})
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        self.assertEqual(wb['Category totals'].max_row, 1)
        self.assertEqual(len(wb['Budget charts']._charts), 1)
        self.assertEqual(len(wb['Budget summary']._images), 1)
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        with ZipFile(io.BytesIO(raw)) as archive:
            chart = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            self.assertEqual(chart.findall('.//c:numCache/c:pt', ns), [])

    def test_proposal_has_written_sections_exact_figures_and_caveats(self):
        original = copy.deepcopy(self.pack)
        with patch.object(server, 'urlopen') as network:
            raw = proposal_pdf_bytes(self.pack)
            network.assert_not_called()
        pdf = PdfReader(io.BytesIO(raw))
        self.assertGreater(len(pdf.pages), 2)
        text = ' '.join(('\n'.join(page.extract_text() for page in pdf.pages)).split())
        for fragment in ['DRAFT EVENT PROPOSAL', '2027-03-20', 'Lecture Theatres', 'RM2,000.00', 'RM240.00', 'RM600.00', 'not a forecast', 'Next actions', 'approvals required', 'No messages are sent']:
            self.assertIn(fragment, text)
        self.assertEqual(self.pack, original)

    def test_xlsx_http_no_cache(self):
        handler = object.__new__(server.Handler)
        handler.wfile = io.BytesIO()
        headers = {}
        handler.send_response = lambda code: None
        handler.send_header = lambda key, value: headers.update({key: value})
        handler.end_headers = lambda: None
        handler.response(200, b'export', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(headers['Cache-Control'], 'no-store')


if __name__ == '__main__':
    unittest.main()
