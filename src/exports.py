"""Local, on-demand exports from a saved Event Pack. No network or mail calls.
The workbook is a factual snapshot, not an editable forecasting model: recorded
figures remain numeric constants and category aggregation uses integer cents.
"""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import re
from xml.sax.saxutils import escape


def recorded_amount(value, cents_key, *rm_keys):
    """Prefer the pack's integer cents; never replace missing evidence with zero."""
    if value.get(cents_key) is not None:
        return float(Decimal(str(value[cents_key])) / 100)
    for key in rm_keys:
        amount = value.get(key)
        if isinstance(amount, (int, float, Decimal)) and not isinstance(amount, bool):
            return float(amount)
    return None


def budget_values(pack):
    b = pack.get("budget") or {}
    return {
        "ceiling": recorded_amount(b, "limitCents", "limit", "limitRM", "budgetRM"),
        "total": recorded_amount(b, "totalCents", "total", "totalRM"),
        "contingency": recorded_amount(b, "contingencyCents", "contingency", "contingencyRM"),
        "gap": recorded_amount(b, "gapCents", "gap", "gapRM"),
        "funds": recorded_amount(b, "organiserFundsCents", "organiserFunds", "organiserFundsRM"),
    }


def text(value):
    if value is None:
        return "Not supplied"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def budget_chart_preview(category_amounts, values):
    """Portable chart image for viewers that do not render native Excel charts.

    Values come from the same snapshot as the native charts. Missing evidence is
    labelled, never drawn as a measured zero. No external fonts/services needed.
    """
    from PIL import Image, ImageDraw, ImageFont
    import math

    def font(size):
        for name in ("DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"):
            try:
                return ImageFont.truetype(name, size)
            except OSError:
                pass
        return ImageFont.load_default(size=size)

    rows = list(category_amounts.items())
    funding = [("Budget ceiling", values["ceiling"]), ("Allocated total", values["total"]),
               ("Protected contingency (included)", values["contingency"]),
               ("Funding gap (unconfirmed)", values["gap"]), ("Organiser funds (recorded)", values["funds"])]
    row_height = 58
    height = 380 + max(1, len(rows)) * row_height + len(funding) * row_height
    image = Image.new("RGB", (1280, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    ink, muted, teal, amber = "#213D42", "#546A68", "#167A70", "#B46516"
    title_font, label_font, small_font = font(29), font(19), font(16)
    draw.text((32, 26), "Estimated allocation by category", font=title_font, fill=ink)
    draw.text((32, 68), "Recorded snapshot in RM - estimates, not vendor quotations", font=small_font, fill=muted)
    left, right = 405, 1000

    def wrap_label(label):
        # Pixel-aware wrapping also handles long words and non-Latin labels.
        lines, line = [], ""
        for character in str(label).replace("\n", " "):
            candidate = line + character
            if draw.textlength(candidate, font=label_font) <= left - 60:
                line = candidate
            else:
                lines.append(line)
                line = character
                if len(lines) == 2:
                    break
        else:
            lines.append(line)
            return "\n".join(lines)
        last = lines[-1]
        while last and draw.textlength(last + "...", font=label_font) > left - 60:
            last = last[:-1]
        lines[-1] = last + "..."
        return "\n".join(lines)

    def amount_label(value, y, color):
        size = 16
        value_font = small_font
        while draw.textlength(value, font=value_font) > 220 and size > 8:
            size -= 1
            value_font = font(size)
        draw.text((1248, y + 4), value, font=value_font, fill=color, anchor="ra")

    def panel(entries, start, include_ceiling=False):
        numeric = [amount for _, amount in entries if amount is not None]
        if include_ceiling and values["ceiling"] is not None:
            numeric.append(values["ceiling"])
        low = min([0] + numeric)
        high = max([0] + numeric)
        step = 10 ** math.floor(math.log10(max(high - low, 1)))
        low = math.floor(low / step) * step
        high = max(low + step, math.ceil(high / step) * step)
        scale = lambda amount: left + (amount - low) / (high - low) * (right - left)
        for index in range(5):
            tick = low + (high - low) * index / 4
            x = scale(tick)
            draw.line((x, start - 10, x, start + len(entries) * row_height - 8), fill="#DFE8E4", width=1)
            draw.text((x, start - 34), f"{tick:,.0f}", font=small_font, fill=muted, anchor="mm")
        for index, (label, amount) in enumerate(entries):
            y = start + index * row_height
            draw.multiline_text((32, y), wrap_label(label), font=label_font, fill=ink, spacing=2)
            if amount is None:
                amount_label("Not supplied", y, muted)
                continue
            color = amber if "contingency" in str(label).lower() or "gap" in str(label).lower() else teal
            x0, x1 = sorted((scale(0), scale(amount)))
            if x1 > x0:
                draw.rectangle((x0, y + 3, x1, y + 29), fill=color)
            amount_label(f"RM{amount:,.2f}", y, ink)

    if rows:
        panel(rows, 145)
    else:
        draw.text((32, 145), "No category allocations supplied", font=label_font, fill=muted)
    funding_top = 205 + max(1, len(rows)) * row_height
    draw.text((32, funding_top), "Budget and funding position", font=title_font, fill=ink)
    panel(funding, funding_top + 95, include_ceiling=True)
    foot = funding_top + 95 + len(funding) * row_height + 15
    draw.text((32, foot), "Contingency is included in allocations. The funding gap is not secured income.", font=small_font, fill=muted)
    draw.text((32, foot + 28), "Static preview. Native Excel charts: 'Budget charts' tab. No approvals or sending implied.", font=small_font, fill=muted)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def budget_xlsx_bytes(pack):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.data_source import AxDataSource, NumData, NumVal, StrData, StrRef, StrVal
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint
    from openpyxl.drawing.image import Image as WorksheetImage
    from openpyxl.styles import Alignment, Font, PatternFill

    def xl_color(css_hex):
        value = css_hex.removeprefix("#").upper()
        if len(value) != 6:
            raise ValueError("Expected six-digit CSS color")
        return "FF" + value

    XL_PRIMARY = xl_color("#167A70")
    XL_LIGHT = xl_color("#E8F3ED")
    XL_INK = xl_color("#213D42")
    XL_WHITE = xl_color("#FFFFFF")
    wb = Workbook()
    wb.properties.title = "EventPilot Budget - Draft Planning Snapshot"
    summary = wb.active
    summary.title = "Budget summary"
    breakdown = wb.create_sheet("Full breakdown")
    categories = wb.create_sheet("Category totals")
    notes = wb.create_sheet("Notes and source")
    money_format = '"RM "#,##0.00;[Red]-"RM "#,##0.00'
    b = pack.get("budget") or {}
    values = budget_values(pack)
    summary.append(["EventPilot budget", "Recorded amount (RM)", "Status / basis"])
    for label, key, note in [
        ("Budget ceiling", "ceiling", "Recorded planning ceiling; not approved spending"),
        ("Allocated total", "total", "Includes contingency; estimates, not vendor quotes"),
        ("Protected contingency", "contingency", "Already included in allocated total; do not add twice"),
        ("Funding gap", "gap", "Unconfirmed gap; prospect asks are not secured income"),
        ("Organiser funds", "funds", b.get("fundingStatus", "Status not supplied")),
    ]:
        summary.append([label, values[key], note])
    summary.append(["Scenario", b.get("scenario", "Not supplied"), "Saved Event Pack; no figures or statuses changed"])
    summary.append(["Snapshot", pack.get("generatedAt", "Not supplied"), "Re-export after the Event Pack changes; no recalculation claims"])
    items = b.get("items", b.get("allocations", []))
    fields = list(dict.fromkeys(key for item in items for key in item))
    breakdown.append(["Recorded allocation (RM)"] + fields)
    category_cents = defaultdict(int)
    missing_categories = set()
    for item in items:
        amount = recorded_amount(item, "amountCents", "amountRM", "totalRM", "amount", "costRM", "total", "cost")
        category = str(item.get("category") or item.get("name") or "Unspecified category")
        if amount is None:
            missing_categories.add(category)
        else:
            category_cents[category] += int(Decimal(str(amount)).scaleb(2).to_integral_value())
        breakdown.append([amount] + [text(item[key]) if isinstance(item.get(key), (dict, list)) else item.get(key) for key in fields])
    categories.append(["Category", "Recorded allocations (RM)", "Budget ceiling (RM)", "Caveat"])
    category_amounts = {}
    for category in dict.fromkeys(str(i.get("category") or i.get("name") or "Unspecified category") for i in items):
        amount = None if category in missing_categories else category_cents[category] / 100
        category_amounts[category] = amount
        categories.append([category, amount, values["ceiling"], "Estimated allocation; not a quote"])

    charts = wb.create_sheet("Budget charts")
    charts.append(["Native Excel charts", "Recorded planning snapshot; estimates, not approved spending"])
    charts.append(["Reading the charts", "Category bars use their own readable scale. Budget ceiling and funding are compared separately below."])
    charts.append(["Compatibility", "If your viewer omits native charts, the opening Budget summary sheet includes a static chart preview."])

    def native_chart(source, first_row, last_row, title, header, height):
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = title
        # openpyxl keeps TextAxis in x_axis and NumericAxis in y_axis even
        # for horizontal bars. Configure by axis type, not visual orientation.
        chart.x_axis.title = "Category / recorded metric"
        chart.y_axis.title = "Malaysian ringgit (RM)"
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"
        chart.add_data(Reference(source, min_col=2, min_row=first_row, max_row=last_row), titles_from_data=True)
        series = chart.series[0]
        labels = [str(source.cell(row, 1).value) for row in range(first_row + 1, last_row + 1)]
        amounts = [source.cell(row, 2).value for row in range(first_row + 1, last_row + 1)]
        # Excel normally regenerates these caches; previewers may require them.
        # Keep references AND complete indexed caches for portable first-open data.
        series.cat = AxDataSource(strRef=StrRef(
            f=str(Reference(source, min_col=1, min_row=first_row + 1, max_row=last_row)),
            strCache=StrData(ptCount=len(labels), pt=[StrVal(idx=i, v=label) for i, label in enumerate(labels)])))
        series.val.numRef.numCache = NumData(formatCode='"RM "#,##0.00', ptCount=len(amounts),
            pt=[NumVal(idx=i, v=amount) for i, amount in enumerate(amounts) if amount is not None])
        series.tx.strRef.strCache = StrData(ptCount=1, pt=[StrVal(idx=0, v=header)])
        series.graphicalProperties.solidFill = "167A70"
        series.graphicalProperties.line.noFill = True
        for index, label in enumerate(labels):
            if "contingency" in label.lower() or "gap" in label.lower():
                point = DataPoint(idx=index)
                point.graphicalProperties.solidFill = "B46516"
                point.graphicalProperties.line.noFill = True
                series.data_points.append(point)
        numeric = [amount for amount in amounts if amount is not None]
        chart.y_axis.scaling.min = min([0] + numeric)
        chart.y_axis.scaling.max = max([1] + numeric) * 1.18
        chart.y_axis.numFmt = '"RM "#,##0'
        chart.legend = None
        chart.height, chart.width = height, 25
        chart.gapWidth = 55
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.dLblPos = "outEnd"
        chart.dataLabels.numFmt = '"RM "#,##0.00'
        chart.display_blanks = "gap"
        chart.visible_cells_only = False
        return chart

    if items:
        category_chart_height = max(10, len(category_amounts) * 0.8 + 4)
        charts.add_chart(native_chart(categories, 1, categories.max_row,
            "Estimated allocation by category", "Recorded allocations (RM)", category_chart_height), "A5")
        funding_anchor = 7 + int(category_chart_height / 0.53)
    else:
        category_chart_height, funding_anchor = 0, 5
    charts.add_chart(native_chart(summary, 1, 6, "Budget ceiling and funding position",
        "Recorded amount (RM)", 10), f"A{funding_anchor}")

    summary.append(["Chart guide", "Preview below", "Native Excel charts are on the Budget charts tab. Missing figures remain unavailable."])
    preview = WorksheetImage(BytesIO(budget_chart_preview(category_amounts, values)))
    preview.width = 960
    preview.height = preview.height * 960 / 1280
    summary.add_image(preview, "A11")
    summary.print_area = f"A1:N{12 + int(preview.height / 20)}"
    charts.print_area = f"A1:N{funding_anchor + 21}"
    notes.append(["Field", "Recorded detail / limitation"])
    notes.append(["Purpose", "Draft planning snapshot for manual review. No approval, booking, sponsorship or email sending is implied."])
    notes.append(["Chart", "Category bars show recorded estimated allocations on a readable scale. A separate chart compares the recorded ceiling, total, included contingency, gap and organiser funds. The opening sheet includes a static image for viewers without native chart support."])
    notes.append(["Totals", "Category totals aggregate recorded integer cents. Snapshot constants, not an editable calculation model."])
    notes.append(["Missing values", "Missing monetary evidence is left blank, never substituted with zero."])
    if pack.get("replay"):
        notes.append(["Replay", text(pack["replay"])])
    for key, value in b.items():
        if key not in ("items", "allocations"):
            notes.append([key, text(value)])
    notes.append(["Validation", text(pack.get("validation"))])
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    # Includes pack-originated field names in header row.
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=XL_PRIMARY)
            cell.font = Font(color=XL_WHITE, bold=True, size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(color=XL_INK, size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    # All pack strings remain text, including leading '='. No formula injection.
                    cell.data_type = "s"
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=XL_LIGHT)
        for col in range(1, ws.max_column + 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col)].width = 35 if col == 1 else 28
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    for row in range(2, 7):
        summary.cell(row, 2).number_format = money_format
    for row in range(2, breakdown.max_row + 1):
        breakdown.cell(row, 1).number_format = money_format
    for row in range(2, categories.max_row + 1):
        for col in (2, 3):
            categories.cell(row, col).number_format = money_format
    summary.column_dimensions["C"].width = 62
    summary.row_dimensions[1].height = 32
    for row in range(2, 10):
        summary.row_dimensions[row].height = 34
    summary.auto_filter.ref = "A1:C8"
    summary["B9"].hyperlink = "#'Budget charts'!A1"
    summary["B9"].font = Font(color=XL_PRIMARY, underline="single", size=11)
    charts.column_dimensions["B"].width = 95
    for row in range(1, 4):
        charts.row_dimensions[row].height = 30
    notes.column_dimensions["B"].width = 110
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def proposal_pdf_bytes(pack):
    """Written draft with exact source data and limitations; no model calls."""
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=46, bottomMargin=46, title="EventPilot Draft Event Proposal", author="EventPilot")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ProposalBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=15, spaceAfter=9, alignment=TA_LEFT, textColor=colors.HexColor("#213D42")))
    story = []

    def paragraph(value, style="ProposalBody"):
        # Normalize unsupported punctuation in built-in fonts without changing figures.
        value = str(value).replace("\u2014", " - ").replace("\u2013", "-").replace("\u2019", "'").replace("\u2192", "to").replace("\u00d7", "x")
        return Paragraph(escape(value).replace("\n", "<br/>"), styles[style])

    def add(value, style="ProposalBody"):
        story.append(paragraph(value, style))

    def heading(value):
        add(value, "Heading2")

    def walk(value, depth=0):
        if isinstance(value, dict):
            facts = []
            for key, item in value.items():
                label = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", str(key)).replace("_", " ").capitalize()
                if isinstance(item, (dict, list)):
                    if facts:
                        add("; ".join(facts) + ".")
                        facts = []
                    if item:
                        add(label + ":")
                        walk(item, depth + 1)
                else:
                    facts.append(label + ": " + text(item))
            if facts:
                add("; ".join(facts) + ".")
        elif isinstance(value, list):
            if value and all(isinstance(item, (int, float)) for item in value):
                add(", ".join(str(item) for item in value))
                return
            for index, item in enumerate(value, 1):
                if isinstance(item, (dict, list)):
                    walk(item, depth + 1)
                else:
                    add(str(index) + ". " + text(item))
        else:
            add(text(value))

    g = pack.get("goal") or {}
    summary = pack.get("summary") or {}
    chosen = (pack.get("dateDecision") or {}).get("chosen") or {}
    proposed_date = chosen.get("date", "Not supplied") if isinstance(chosen, dict) else str(chosen)
    venues = pack.get("venues") or {}
    selected = venues.get("selected") or venues.get("recommended") or {}
    values = budget_values(pack)
    money = lambda amount: "Not supplied" if amount is None else f"RM{amount:,.2f}"
    add("DRAFT EVENT PROPOSAL", "Title")
    add("For organiser review and manual submission only. This download does not send anything or confirm any approval, booking, funding or participation.")
    add("Source Event Pack generated: " + str(pack.get("generatedAt", "Not supplied")))
    if pack.get("replay"):
        add("Source is a recorded replay, not new research: " + text(pack["replay"]))
    heading("1. Event summary and purpose")
    add(summary.get("narrative", text(summary)) if isinstance(summary, dict) else text(summary))
    add(f"The proposed event is {g.get('eventType', 'not supplied')} at {g.get('campus', 'a campus not supplied')}. Its stated purpose is {g.get('cause', 'not supplied')}. The recorded attendance target is {g.get('attendance', 'not supplied')}; this is a planning target, not confirmed attendance.")
    heading("2. Proposed date and calendar checks")
    add(f"The chosen planning date is {proposed_date}. It is proposed, not a confirmed calendar clearance. Campus and programme-specific examination dates require confirmation.")
    walk({k: v for k, v in chosen.items() if k != "weather"} if isinstance(chosen, dict) else chosen)
    decision = pack.get("dateDecision") or {}
    for backup in decision.get("backups", []):
        if isinstance(backup, dict):
            walk({k: v for k, v in backup.items() if k not in ("weather", "sources", "sourceUrls")})
        else:
            add("Backup date: " + str(backup))
    walk({k: v for k, v in decision.items() if k not in ("chosen", "candidates", "backups", "calendar")})
    heading("3. Proposed venue and approval requirements")
    add(f"The selected proposal is {selected.get('name', 'not supplied')}. Recorded approval status: {selected.get('approvalStatus', venues.get('approvalStatus', 'not supplied'))}. A cited capacity or suitability assessment is not a booking or safety approval.")
    walk(selected)
    walk({k: v for k, v in venues.items() if k not in ("selected", "options", "recommended")})
    heading("4. Budget and funding position (RM)")
    add(f"Recorded ceiling: {money(values['ceiling'])}; allocated total: {money(values['total'])}; protected contingency: {money(values['contingency'])}; funding gap: {money(values['gap'])}; organiser funds: {money(values['funds'])}. Contingency is already included in the allocation total. All allocations are estimates, not vendor quotations, and prospective sponsor asks are not secured income.")
    rows = [[paragraph("Category"), paragraph("Recorded RM"), paragraph("Status / basis")]]
    for item in (pack.get("budget") or {}).get("items", []):
        rows.append([paragraph(item.get("category", "Not supplied")), paragraph(money(recorded_amount(item, "amountCents", "amountRM"))), paragraph(item.get("status", "Not supplied"))])
    table = Table(rows, colWidths=[190, 80, 241], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F3ED")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), .4, colors.HexColor("#D9E3DA")), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7)]))
    story.append(table)
    story.append(Spacer(1, 12))
    walk({k: v for k, v in (pack.get("budget") or {}).items() if k not in ("items", "allocations")})
    heading("5. Weather limitations and risk plan")
    weather = pack.get("weather") or {}
    add("Recorded weather evidence mode: " + str(weather.get("mode", "Not supplied")) + ". Historical climatology is not an event-date forecast; a source forecast is not a guarantee. Missing evidence does not imply low risk.")
    walk({k: v for k, v in weather.items() if k not in ("samples", "sourceUrls")})
    walk(pack.get("risks", "Not supplied"))
    heading("6. Programme and operational outline")
    walk(pack.get("programme", "Not supplied"))
    heading("7. Next actions before approval or submission")
    walk(summary.get("nextActions", "Not supplied") if isinstance(summary, dict) else "Not supplied")
    walk(pack.get("rechecks", "Not supplied"))
    heading("8. Validation and outstanding confirmations")
    walk(pack.get("validation", "Not supplied"))
    add("Passing planning checks is not real-world approval. Review this draft, confirm facts with the responsible campus teams and obtain permission before submitting, spending, publishing or contacting anyone. No messages are sent by this feature.")

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#546A68"))
        canvas.drawString(42, 26, "EventPilot | DRAFT - approvals required | Never sends")
        canvas.drawRightString(A4[0] - 42, 26, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
